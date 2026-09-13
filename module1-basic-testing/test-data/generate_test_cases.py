from __future__ import annotations

import argparse
import shutil
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment


def make_case(
    case_id: str,
    item: str,
    title: str,
    criticality: str,
    precondition: str,
    test_input: str,
    procedure: str,
    expected: str,
    method: str,
    remark: str = "",
) -> dict[str, str]:
    return {
        "id": case_id,
        "item": item,
        "title": title,
        "criticality": criticality,
        "automation": "Y",
        "requirement_change": "N",
        "precondition": precondition,
        "input": test_input,
        "procedure": procedure,
        "expected": expected,
        "actual": "与预期一致，自动化测试通过",
        "status": "OK",
        "method": method,
        "remark": remark,
    }


def cases() -> list[dict[str, str]]:
    cases = [
        make_case("A-001", "认证-JWT", "签发并解析有效 JWT", "High", "JwtService 使用 test-secret，系统时间有效", "userId=7，username=alice，role=ADMIN", "调用 issue 生成 token，再调用 parse", "解析成功且三项用户声明与输入一致", "等价类"),
        make_case("A-002", "认证-JWT", "签名被修改的 JWT 被拒绝", "High", "已获得一个有效 JWT", "将 token 最后一位替换为其他字符", "调用 parse 解析篡改 token", "返回空 Optional，不建立身份", "等价类"),
        make_case("A-003", "认证-JWT", "已过期 JWT 被拒绝", "High", "JwtService 过期时间设置为 -1 分钟", "刚签发的 token", "调用 parse", "返回空 Optional", "边界值"),
        make_case("A-004", "认证-JWT", "空字符串 token 被拒绝", "High", "JwtService 可用", "token=空字符串", "调用 parse", "返回空 Optional", "等价类"),
        make_case("A-005", "认证-JWT", "无分隔符 token 被拒绝", "High", "JwtService 可用", "token=abc", "调用 parse", "返回空 Optional", "等价类"),
        make_case("A-006", "认证-JWT", "三段式但签名结构错误的 token 被拒绝", "High", "JwtService 可用", "token=a.b.c", "调用 parse", "返回空 Optional", "等价类"),
        make_case("A-007", "认证-JWT", "非法 Base64 token 被拒绝", "High", "JwtService 可用", "token=%%%.___", "调用 parse", "返回空 Optional", "等价类"),
        make_case("A-008", "认证-JWT", "末尾分隔符 token 被拒绝", "High", "JwtService 可用", "token=a.", "调用 parse", "返回空 Optional", "等价类"),
        make_case("A-009", "认证-过滤器", "有效 Bearer token 建立管理员身份", "High", "JwtAuthenticationFilter 使用有效 JwtService", "Authorization=Bearer + 有效 ADMIN token", "通过过滤器执行请求", "SecurityContext 中 principal 为 JwtUser，权限为 ROLE_ADMIN，且继续过滤链", "场景法"),
        make_case("A-010", "认证-过滤器", "过期 Bearer token 不建立身份", "High", "过滤器使用过期 JwtService", "Authorization=Bearer + 过期 token", "通过过滤器执行请求", "SecurityContext 无认证，但请求继续过滤链", "等价类"),
        make_case("A-011", "认证-过滤器", "非 Bearer Authorization 被忽略", "Medium", "过滤器可用", "Authorization=Basic abc", "通过过滤器执行请求", "不解析 token，SecurityContext 无认证，请求继续", "等价类"),
        make_case("A-012", "认证-过滤器", "格式错误 Bearer token 被忽略", "High", "过滤器可用", "Authorization=Bearer malformed-token", "通过过滤器执行请求", "不建立身份，请求继续过滤链", "等价类"),
        make_case("A-013", "认证-登录", "正确用户名密码登录成功", "High", "数据库返回 alice，密码哈希匹配", "alice/correct-password", "调用 AuthService.login", "返回用户信息和 JWT，角色与组织信息正确", "场景法"),
        make_case("A-014", "认证-登录", "不存在用户名登录失败", "High", "数据库查不到 unknown", "unknown/password", "调用 AuthService.login", "抛出统一的用户名或密码错误，不签发 JWT", "等价类"),
        make_case("A-015", "认证-登录", "错误密码登录失败", "High", "数据库存在 alice，密码哈希不匹配", "alice/wrong-password", "调用 AuthService.login", "抛出统一的用户名或密码错误，不签发 JWT", "等价类"),
        make_case("A-016", "访问控制", "管理员可创建任意运单", "High", "用户角色为 ADMIN", "任意组织运单记录", "调用 requireCanCreateWaybill", "不抛出 AccessDeniedException", "等价类"),
        make_case("A-017", "访问控制", "托运人可创建本组织运单", "High", "用户角色 CONSIGNOR，组织 10；运单托运人组织 10", "组织匹配的运单", "调用 requireCanCreateWaybill", "不抛出 AccessDeniedException", "等价类"),
        make_case("A-018", "访问控制", "托运人不能创建他组织运单", "High", "用户角色 CONSIGNOR，组织 99；运单托运人组织 10", "组织不匹配的运单", "调用 requireCanCreateWaybill", "抛出 AccessDeniedException", "等价类"),
        make_case("A-019", "访问控制", "跨链网关不能查看普通运单", "High", "用户角色 CROSS_CHAIN_GATEWAY，待过滤运单列表非空", "包含普通运单的列表", "调用 filterVisibleWaybills", "返回空列表", "等价类"),
        make_case("A-020", "访问控制", "上传者可下载自己的文件", "High", "文件 uploaderId 与用户 id 相同", "uploaderId=7，userId=7", "调用 canDownload", "返回 true", "等价类"),
        make_case("A-021", "访问控制", "无有效授权的他人不能下载文件", "High", "用户不是上传者，用户/组织/角色均无 ACTIVE 授权", "其他用户下载文件", "调用 canDownload", "返回 false", "等价类"),
        make_case("A-022", "访问控制", "授权目标角色按白名单校验", "Medium", "访问控制服务可用", "CONSIGNEE、CUSTOMS、FORWARDER、CARRIER", "逐个调用 requireAllowedGrantTarget", "前三者通过，CARRIER 被拒绝", "等价类"),
        make_case("A-023", "访问控制", "跨链操作角色校验", "High", "访问控制服务可用", "ADMIN、CROSS_CHAIN_GATEWAY、CONSIGNOR", "逐个调用 requireCrossChainOperator", "前两者通过，CONSIGNOR 被拒绝", "等价类"),
        make_case("B-001", "跨链-前置检查", "源链与目标链相同被拒绝", "High", "跨链操作用户已认证", "source=CHINA，target=CHINA", "调用 precheck", "结果 passed=false，包含源链和目标链失败项", "边界值"),
        make_case("B-002", "跨链-前置检查", "已过期共享请求被拒绝", "High", "运单、文件、源链记录均存在", "expiresAt=当前时间之前", "调用 precheck", "结果 passed=false，有效期合法项失败", "边界值"),
        make_case("B-003", "跨链-前置检查", "依赖完整且有效期未来时检查通过", "High", "MySQL 运单、文件摘要、源链存证均存在，源链匹配", "source=CHINA，target=EUROPE，expiresAt=未来", "调用 precheck", "8 个检查项全部通过", "场景法"),
        make_case("B-004", "跨链-任务创建", "前置检查通过后创建 CREATED 任务", "High", "跨链前置检查全部通过，当前用户 id=7", "WB-001，CHINA -> EUROPE，有效期未来", "调用 create", "保存带 CCT-/XCC- 标识、createdBy=7、状态 CREATED 的任务", "场景法"),
        make_case("B-005", "跨链-任务创建", "前置检查失败时不创建任务", "High", "运单存在但文件和源链记录缺失", "WB-001，CHINA -> EUROPE", "调用 create", "抛出前置检查未通过异常，仓储不保存任务", "等价类"),
        make_case("B-006", "跨链-执行", "跨链成功执行并确认目标链回执", "High", "任务为 CREATED，源链记录存在，目标链 Mock 返回成功回执", "目标 tx=europe-tx-001", "调用 execute", "任务最终为 CONFIRMED，生成 credential/message/receipt/evidence 哈希", "场景法"),
        make_case("B-007", "跨链-retry", "CONFIRMED 任务 retry 不重复写目标链", "High", "任务已成功执行并有完整哈希", "再次调用 execute 同一 taskId", "连续调用两次 execute", "第二次直接返回确认结果，Fabric 接收方法总共只调用一次", "场景法"),
        make_case("B-008", "跨链-失败处理", "源链记录缺失时任务转 FAILED", "High", "任务存在但 Fabric 源链查询为空", "CCT-001", "调用 execute", "任务状态为 FAILED，保存错误信息，不调用目标链接收", "等价类"),
        make_case("B-009", "跨链-核验", "未确认任务不能执行跨链核验", "High", "任务状态为 CREATED", "CCT-001", "调用 verify", "返回 passed=false，记录核验时间和失败原因", "状态转换"),
    ]
    if len(cases) != 32:
        raise AssertionError(f"expected 32 active cases, got {len(cases)}")
    return cases


def build_workbook(template: Path, output: Path) -> None:
    workbook = openpyxl.load_workbook(template)
    information = workbook.worksheets[0]
    test_sheet = workbook.worksheets[1]
    data = cases()

    information["E7"] = "1.0"
    information["J7"] = "课程作业-模块一"
    information["E8"] = "Software-Testing-Practice--Gray-Box-Group"
    information["J8"] = "STP-GRAY-BOX-2026"
    information["E9"] = "ZhuJiacheng"
    information["J9"] = "2026.09.13"
    information["E10"] = "待复核"
    information["E11"] = "待审批"
    information["E13"] = f"=COUNTA('{test_sheet.title}'!A2:A{len(data) + 1})"
    information["B16"] = "本批次采用传统测试方法，当前范围为模块 A 认证与访问控制、模块 B 跨链任务状态控制。"
    information["B20"] = "2026.09.13"
    information["C20"] = "1.00"
    information["D20"] = "STP-INIT"
    information["E20"] = "建立 A+B 测试基线"
    information["K20"] = "ZhuJiacheng"

    for row in range(2, 2 + len(data)):
        for column in range(1, 14):
            cell = test_sheet.cell(row=row, column=column)
            cell._style = copy(test_sheet.cell(row=2, column=column)._style)
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

    for row, test_case in enumerate(data, start=2):
        values = [
            test_case["id"],
            test_case["item"],
            test_case["title"],
            test_case["criticality"],
            test_case["automation"],
            test_case["requirement_change"],
            test_case["precondition"],
            test_case["input"],
            test_case["procedure"],
            test_case["expected"],
            test_case["actual"],
            test_case["status"],
            test_case["method"] + (f"；{test_case['remark']}" if test_case["remark"] else ""),
        ]
        for column, value in enumerate(values, start=1):
            test_sheet.cell(row=row, column=column).value = value
        test_sheet.row_dimensions[row].height = 58

    test_sheet.freeze_panes = "A2"
    test_sheet.auto_filter.ref = f"A1:M{len(data) + 1}"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)


def main() -> None:
    parser = argparse.ArgumentParser(description="Create a module-one test-case workbook from the course template.")
    parser.add_argument("--template", type=Path, required=True)
    parser.add_argument("--source-output", type=Path, required=True)
    parser.add_argument("--repo-output", type=Path, required=True)
    args = parser.parse_args()

    build_workbook(args.template, args.source_output)
    args.repo_output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(args.source_output, args.repo_output)
    print(f"created: {args.source_output}")
    print(f"copied:  {args.repo_output}")
    print(f"cases:   {len(cases())}")


if __name__ == "__main__":
    main()
